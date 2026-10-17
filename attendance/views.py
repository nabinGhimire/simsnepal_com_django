from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Q, F
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from datetime import date, datetime, timedelta
from itertools import groupby

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from sms.models import Student, SchoolGrade, Section, StudentSession, Subject, SchoolBranch, EduSession, BranchUser, GradeLevel
from .models import DailyAttendance, AttendanceAuditLog
import nepali_datetime

NEPALI_MONTHS = [
    (1, "Baisakh (वैशाख)"),
    (2, "Jestha (जेठ)"),
    (3, "Asar (असार)"),
    (4, "Shrawan (साउन)"),
    (5, "Bhadra (भदौ)"),
    (6, "Ashwin (असोज)"),
    (7, "Kartik (कात्तिक)"),
    (8, "Mangsir (मंसिर)"),
    (9, "Poush (पुस)"),
    (10, "Magh (माघ)"),
    (11, "Falgun (फागुन)"),
    (12, "Chaitra (चैत)"),
]

NEPALI_MONTH_SHORT_NAMES = [
    "Baisakh", "Jestha", "Asar", "Shrawan", "Bhadra", "Ashwin",
    "Kartik", "Mangsir", "Poush", "Magh", "Falgun", "Chaitra"
]

def get_nepali_month_ad_range(year_bs, month_bs):
    try:
        start_bs = nepali_datetime.date(year_bs, month_bs, 1)
        start_ad = start_bs.to_datetime_date()
        if month_bs == 12:
            next_start_bs = nepali_datetime.date(year_bs + 1, 1, 1)
        else:
            next_start_bs = nepali_datetime.date(year_bs, month_bs + 1, 1)
        end_ad = next_start_bs.to_datetime_date() - timedelta(days=1)
        return start_ad, end_ad
    except Exception:
        today_bs = nepali_datetime.date.today()
        start_bs = nepali_datetime.date(today_bs.year, today_bs.month, 1)
        return start_bs.to_datetime_date(), date.today()


def parse_bs_or_ad_date(date_str):
    if not date_str:
        bs_today = nepali_datetime.date.today()
        return bs_today.to_datetime_date(), bs_today, f"{bs_today.year}-{bs_today.month:02d}-{bs_today.day:02d}"
    try:
        parts = [int(p) for p in date_str.split('-')]
        if len(parts) == 3:
            if 2050 <= parts[0] <= 2120:
                bs_obj = nepali_datetime.date(parts[0], parts[1], parts[2])
                ad_obj = bs_obj.to_datetime_date()
                return ad_obj, bs_obj, f"{bs_obj.year}-{bs_obj.month:02d}-{bs_obj.day:02d}"
            else:
                ad_obj = date(parts[0], parts[1], parts[2])
                bs_obj = nepali_datetime.date.from_datetime(ad_obj)
                return ad_obj, bs_obj, f"{bs_obj.year}-{bs_obj.month:02d}-{bs_obj.day:02d}"
    except Exception:
        pass

    bs_today = nepali_datetime.date.today()
    return bs_today.to_datetime_date(), bs_today, f"{bs_today.year}-{bs_today.month:02d}-{bs_today.day:02d}"


def get_current_school(request):
    sso_business = request.session.get('sso_business', {})
    business_id = sso_business.get('id')
    if business_id:
        branch = SchoolBranch.objects.filter(shortcode=business_id).first()
        if branch:
            return branch
    branchuser = BranchUser.objects.filter(user=request.user).first()
    if branchuser:
        return branchuser.school
    return SchoolBranch.objects.first()


def get_current_session(request):
    active_session_id = request.session.get('active_session_id')
    if active_session_id:
        sess = EduSession.objects.filter(id=active_session_id).first()
        if sess:
            return sess
    return EduSession.objects.filter(status=True).order_by('-year').first()


@login_required
def attendance_dashboard(request):
    school = get_current_school(request)
    current_session = get_current_session(request)
    
    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')
    selected_subject_id = request.GET.get('subject')
    selected_date_raw = request.GET.get('date')
    mode = request.GET.get('mode', 'view')

    selected_date, selected_bs_date_obj, selected_bs_date_str = parse_bs_or_ad_date(selected_date_raw)

    selected_grade = None
    selected_section = None
    selected_subject = None
    sections = []
    subjects = []
    student_sessions = []
    attendance_map = {}
    taken_by_user = None
    is_already_taken = False
    audit_logs = []

    if selected_grade_id:
        selected_grade = grades.filter(id=selected_grade_id).first()
        if not selected_grade:
            messages.error(request, f"Grade not found. The selected grade (ID: {selected_grade_id}) does not exist or is not active.")
            return redirect('attendance:dashboard')
        sections = Section.objects.filter(grade=selected_grade)
        subjects = Subject.objects.filter(grade=selected_grade)

        if selected_section_id:
            selected_section = get_object_or_404(sections, id=selected_section_id)
        if selected_subject_id:
            selected_subject = get_object_or_404(subjects, id=selected_subject_id)

        ss_query = StudentSession.objects.filter(
            grade=selected_grade,
            session=current_session,
            status=True
        ).select_related('student', 'section')

        if selected_section:
            ss_query = ss_query.filter(section=selected_section)

        student_sessions = list(ss_query.order_by('section__section', 'roll_no', 'student__name'))

        att_query = DailyAttendance.objects.filter(
            grade=selected_grade,
            date=selected_date,
            session=current_session
        )
        if selected_section:
            att_query = att_query.filter(section=selected_section)
        if selected_subject:
            att_query = att_query.filter(subject=selected_subject)
        else:
            att_query = att_query.filter(subject__isnull=True)

        is_already_taken = att_query.exists()

        for record in att_query:
            attendance_map[record.reg_no_id] = record
            if not taken_by_user and record.taken_by:
                taken_by_user = record.taken_by

        # Fetch Audit Logs if editing existing attendance
        if is_already_taken:
            audit_logs = AttendanceAuditLog.objects.filter(
                attendance__in=att_query
            ).select_related('updated_by', 'attendance', 'attendance__reg_no')[:20]

        # Handle POST: Save/Update attendance with Audit Logging
        if request.method == 'POST':
            updated_count = 0
            for ss in student_sessions:
                student = ss.student
                status_val = request.POST.get(f'attendance_{student.reg_no}')
                is_present = (status_val == 'present')
                absent_informed = (request.POST.get(f'informed_{student.reg_no}') == 'yes')
                remarks = request.POST.get(f'remarks_{student.reg_no}', '')

                # Check existing record for audit logging
                existing = DailyAttendance.objects.filter(
                    reg_no=student,
                    date=selected_date,
                    subject=selected_subject
                ).first()

                rec, created = DailyAttendance.objects.update_or_create(
                    reg_no=student,
                    date=selected_date,
                    subject=selected_subject,
                    defaults={
                        'school': school or student.school,
                        'session': current_session,
                        'grade': selected_grade,
                        'section': ss.section,
                        'nepali_date': selected_bs_date_obj,
                        'is_present': is_present,
                        'absent_informed': absent_informed,
                        'remarks': remarks,
                        'taken_by': request.user
                    }
                )

                if not created and existing:
                    # Log changes if status or remarks altered
                    if existing.is_present != is_present or existing.remarks != remarks or existing.absent_informed != absent_informed:
                        AttendanceAuditLog.objects.create(
                            attendance=rec,
                            updated_by=request.user,
                            previous_status=existing.is_present,
                            new_status=is_present,
                            previous_remarks=existing.remarks,
                            new_remarks=remarks
                        )
                        updated_count += 1

            if updated_count > 0:
                messages.success(request, f"Attendance for Grade {selected_grade.grade_name} updated successfully ({updated_count} changes logged in Audit Trail).")
            else:
                messages.success(request, f"Attendance for Grade {selected_grade.grade_name} saved successfully.")

            return redirect(f"{request.path}?grade={selected_grade.id}&section={selected_section_id or ''}&subject={selected_subject_id or ''}&date={selected_bs_date_str}&mode=view")

    cards_data = []
    cards_by_level = []
    overall_summary = {
        'total_enrolled': 0,
        'evaluated_students': 0,
        'pending_students': 0,
        'present_students': 0,
        'absent_students': 0,
        'informed_absent': 0,
        'uninformed_absent': 0,
        'total_boys': 0,
        'evaluated_boys': 0,
        'present_boys': 0,
        'informed_absent_boys': 0,
        'uninformed_absent_boys': 0,
        'absent_boys': 0,
        'total_girls': 0,
        'evaluated_girls': 0,
        'present_girls': 0,
        'informed_absent_girls': 0,
        'uninformed_absent_girls': 0,
        'absent_girls': 0,
        'sections_taken': 0,
        'total_sections': 0
    }

    chart_class_names = []
    chart_present_rates = []
    chart_absent_rates = []
    chart_boy_present_rates = []
    chart_boy_absent_rates = []
    chart_girl_present_rates = []
    chart_girl_absent_rates = []

    if not selected_grade:
        all_active_ss = StudentSession.objects.filter(
            session=current_session,
            grade__in=grades,
            status=True
        ).select_related('grade', 'section', 'student', 'grade__level')

        all_today_att = DailyAttendance.objects.filter(
            date=selected_date,
            session=current_session
        ).select_related('reg_no', 'taken_by', 'subject')

        for grade in grades:
            grade_sections = list(Section.objects.filter(grade=grade))
            if grade_sections:
                for sec in grade_sections:
                    sec_stds = [s for s in all_active_ss if s.grade_id == grade.id and s.section_id == sec.id]
                    if not sec_stds:
                        continue
                    sec_atts = [a for a in all_today_att if a.grade_id == grade.id and a.section_id == sec.id]
                    card = build_card_dict(grade, sec, sec_stds, sec_atts)
                    cards_data.append(card)

                    chart_class_names.append(f"{grade.grade_name}-{sec.section}")
                    chart_present_rates.append(card['percentage'] if card['is_taken'] else 0)
                    absent_pct = round((card['absent_total'] / card['total_students'] * 100), 1) if (card['is_taken'] and card['total_students'] > 0) else 0
                    chart_absent_rates.append(absent_pct)

                    b_pres = round((card['present_boys'] / card['total_boys'] * 100), 1) if (card['is_taken'] and card['total_boys'] > 0) else 0
                    b_abs = round((card['absent_boys'] / card['total_boys'] * 100), 1) if (card['is_taken'] and card['total_boys'] > 0) else 0
                    g_pres = round((card['present_girls'] / card['total_girls'] * 100), 1) if (card['is_taken'] and card['total_girls'] > 0) else 0
                    g_abs = round((card['absent_girls'] / card['total_girls'] * 100), 1) if (card['is_taken'] and card['total_girls'] > 0) else 0

                    chart_boy_present_rates.append(b_pres)
                    chart_boy_absent_rates.append(b_abs)
                    chart_girl_present_rates.append(g_pres)
                    chart_girl_absent_rates.append(g_abs)
            else:
                sec_stds = [s for s in all_active_ss if s.grade_id == grade.id]
                if not sec_stds:
                    continue
                sec_atts = [a for a in all_today_att if a.grade_id == grade.id]
                card = build_card_dict(grade, None, sec_stds, sec_atts)
                cards_data.append(card)

                chart_class_names.append(f"{grade.grade_name}")
                chart_present_rates.append(card['percentage'] if card['is_taken'] else 0)
                absent_pct = round((card['absent_total'] / card['total_students'] * 100), 1) if (card['is_taken'] and card['total_students'] > 0) else 0
                chart_absent_rates.append(absent_pct)

                b_pres = round((card['present_boys'] / card['total_boys'] * 100), 1) if (card['is_taken'] and card['total_boys'] > 0) else 0
                b_abs = round((card['absent_boys'] / card['total_boys'] * 100), 1) if (card['is_taken'] and card['total_boys'] > 0) else 0
                g_pres = round((card['present_girls'] / card['total_girls'] * 100), 1) if (card['is_taken'] and card['total_girls'] > 0) else 0
                g_abs = round((card['absent_girls'] / card['total_girls'] * 100), 1) if (card['is_taken'] and card['total_girls'] > 0) else 0

                chart_boy_present_rates.append(b_pres)
                chart_boy_absent_rates.append(b_abs)
                chart_girl_present_rates.append(g_pres)
                chart_girl_absent_rates.append(g_abs)

        for level, group in groupby(cards_data, key=lambda c: c['grade'].level):
            cards_by_level.append({
                'level': level,
                'cards': list(group)
            })

        for c in cards_data:
            overall_summary['total_sections'] += 1
            overall_summary['total_enrolled'] += c['total_students']
            overall_summary['total_boys'] += c['total_boys']
            overall_summary['total_girls'] += c['total_girls']

            if c['is_taken']:
                overall_summary['sections_taken'] += 1
                overall_summary['evaluated_students'] += c['total_students']
                overall_summary['present_students'] += c['present_total']
                overall_summary['informed_absent'] += c['informed_absent']
                overall_summary['uninformed_absent'] += c['uninformed_absent']
                overall_summary['absent_students'] += c['absent_total']

                overall_summary['evaluated_boys'] += c['total_boys']
                overall_summary['present_boys'] += c['present_boys']
                overall_summary['informed_absent_boys'] += c['informed_absent_boys']
                overall_summary['uninformed_absent_boys'] += c['uninformed_absent_boys']
                overall_summary['absent_boys'] += c['absent_boys']

                overall_summary['evaluated_girls'] += c['total_girls']
                overall_summary['present_girls'] += c['present_girls']
                overall_summary['informed_absent_girls'] += c['informed_absent_girls']
                overall_summary['uninformed_absent_girls'] += c['uninformed_absent_girls']
                overall_summary['absent_girls'] += c['absent_girls']

        overall_summary['pending_students'] = overall_summary['total_enrolled'] - overall_summary['evaluated_students']

        if overall_summary['evaluated_students'] > 0:
            overall_summary['evaluated_percentage'] = round((overall_summary['present_students'] / overall_summary['evaluated_students']) * 100, 1)
        else:
            overall_summary['evaluated_percentage'] = 0

        if overall_summary['total_enrolled'] > 0:
            overall_summary['total_school_percentage'] = round((overall_summary['present_students'] / overall_summary['total_enrolled']) * 100, 1)
        else:
            overall_summary['total_school_percentage'] = 0

        if overall_summary['evaluated_boys'] > 0:
            overall_summary['boy_present_pct'] = round((overall_summary['present_boys'] / overall_summary['evaluated_boys']) * 100, 1)
            overall_summary['boy_absent_pct'] = round((overall_summary['absent_boys'] / overall_summary['evaluated_boys']) * 100, 1)
        else:
            overall_summary['boy_present_pct'] = 0
            overall_summary['boy_absent_pct'] = 0

        if overall_summary['evaluated_girls'] > 0:
            overall_summary['girl_present_pct'] = round((overall_summary['present_girls'] / overall_summary['evaluated_girls']) * 100, 1)
            overall_summary['girl_absent_pct'] = round((overall_summary['absent_girls'] / overall_summary['evaluated_girls']) * 100, 1)
        else:
            overall_summary['girl_present_pct'] = 0
            overall_summary['girl_absent_pct'] = 0

        if overall_summary['evaluated_boys'] > 0:
            overall_summary['informed_boy_pct'] = round((overall_summary['informed_absent_boys'] / overall_summary['evaluated_boys']) * 100, 1)
            overall_summary['uninformed_boy_pct'] = round((overall_summary['uninformed_absent_boys'] / overall_summary['evaluated_boys']) * 100, 1)
        else:
            overall_summary['informed_boy_pct'] = 0
            overall_summary['uninformed_boy_pct'] = 0

        if overall_summary['evaluated_girls'] > 0:
            overall_summary['informed_girl_pct'] = round((overall_summary['informed_absent_girls'] / overall_summary['evaluated_girls']) * 100, 1)
            overall_summary['uninformed_girl_pct'] = round((overall_summary['uninformed_absent_girls'] / overall_summary['evaluated_girls']) * 100, 1)
        else:
            overall_summary['informed_girl_pct'] = 0
            overall_summary['uninformed_girl_pct'] = 0

        overall_summary['pending_boys'] = overall_summary['evaluated_boys'] - overall_summary['present_boys'] - overall_summary['absent_boys'] if overall_summary['evaluated_boys'] > overall_summary['present_boys'] + overall_summary['absent_boys'] else 0
        overall_summary['pending_girls'] = overall_summary['evaluated_girls'] - overall_summary['present_girls'] - overall_summary['absent_girls'] if overall_summary['evaluated_girls'] > overall_summary['present_girls'] + overall_summary['absent_girls'] else 0

        if overall_summary['total_sections'] > 0:
            overall_summary['recorded_pct'] = round((overall_summary['sections_taken'] / overall_summary['total_sections']) * 100, 1)
            overall_summary['unrecorded_pct'] = round(((overall_summary['total_sections'] - overall_summary['sections_taken']) / overall_summary['total_sections']) * 100, 1)
        else:
            overall_summary['recorded_pct'] = 0
            overall_summary['unrecorded_pct'] = 0

        overall_summary['unrecorded_sections'] = overall_summary['total_sections'] - overall_summary['sections_taken']

        if overall_summary['evaluated_students'] > 0:
            overall_summary['absent_pct'] = round((overall_summary['absent_students'] / overall_summary['evaluated_students']) * 100, 1)
        else:
            overall_summary['absent_pct'] = 0

        if overall_summary['absent_students'] > 0:
            overall_summary['informed_pct'] = round((overall_summary['informed_absent'] / overall_summary['absent_students']) * 100, 1)
            overall_summary['uninformed_pct'] = round((overall_summary['uninformed_absent'] / overall_summary['absent_students']) * 100, 1)
        else:
            overall_summary['informed_pct'] = 0
            overall_summary['uninformed_pct'] = 0

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'selected_grade': selected_grade,
        'sections': sections,
        'selected_section': selected_section,
        'subjects': subjects,
        'selected_subject': selected_subject,
        'selected_date': selected_date.isoformat(),
        'selected_bs_date_str': selected_bs_date_str,
        'mode': mode,
        'student_sessions': student_sessions,
        'attendance_map': attendance_map,
        'taken_by_user': taken_by_user,
        'is_already_taken': is_already_taken,
        'audit_logs': audit_logs,
        'cards_data': cards_data,
        'cards_by_level': cards_by_level,
        'overall_summary': overall_summary,
        'chart_class_names': chart_class_names,
        'chart_present_rates': chart_present_rates,
        'chart_absent_rates': chart_absent_rates,
        'chart_boy_present_rates': chart_boy_present_rates,
        'chart_boy_absent_rates': chart_boy_absent_rates,
        'chart_girl_present_rates': chart_girl_present_rates,
        'chart_girl_absent_rates': chart_girl_absent_rates,
    }

    # Grade-level summary stats for the top cards
    grade_summary = {
        'total': len(student_sessions),
        'present': 0,
        'absent': 0,
        'informed_absent': 0,
        'uninformed_absent': 0,
        'boys_total': 0,
        'boys_present': 0,
        'girls_total': 0,
        'girls_present': 0,
    }
    for ss in student_sessions:
        if ss.student.gender:
            grade_summary['boys_total'] += 1
        else:
            grade_summary['girls_total'] += 1
    for reg_no, rec in attendance_map.items():
        if rec.is_present:
            grade_summary['present'] += 1
            # find the student to count gender
            for ss in student_sessions:
                if ss.student.reg_no == reg_no:
                    if ss.student.gender:
                        grade_summary['boys_present'] += 1
                    else:
                        grade_summary['girls_present'] += 1
                    break
        else:
            grade_summary['absent'] += 1
            if rec.absent_informed:
                grade_summary['informed_absent'] += 1
            else:
                grade_summary['uninformed_absent'] += 1
    grade_summary['boys_absent'] = grade_summary['boys_total'] - grade_summary['boys_present']
    grade_summary['girls_absent'] = grade_summary['girls_total'] - grade_summary['girls_present']
    if grade_summary['total'] > 0:
        grade_summary['percentage'] = round((grade_summary['present'] / grade_summary['total']) * 100, 1)
    else:
        grade_summary['percentage'] = 0

    context['grade_summary'] = grade_summary
    
    return render(request, 'attendance/dashboard.html', context)


@login_required
def transfer_attendance_date(request):
    """
    Admin Date Transfer Tool: Move attendance from one BS date to another.
    """
    school = get_current_school(request)
    current_session = get_current_session(request)

    if request.method == 'POST':
        grade_id = request.POST.get('grade')
        section_id = request.POST.get('section')
        from_bs_date = request.POST.get('from_date')
        to_bs_date = request.POST.get('to_date')

        from_ad, from_bs_obj, _ = parse_bs_or_ad_date(from_bs_date)
        to_ad, to_bs_obj, _ = parse_bs_or_ad_date(to_bs_date)

        records = DailyAttendance.objects.filter(
            session=current_session,
            date=from_ad
        )
        if grade_id:
            records = records.filter(grade_id=grade_id)
        if section_id:
            records = records.filter(section_id=section_id)

        count = records.count()
        if count > 0:
            records.update(
                date=to_ad,
                nepali_date=to_bs_obj
            )
            messages.success(request, f"Successfully transferred {count} attendance records from BS {from_bs_date} to BS {to_bs_date}.")
        else:
            messages.warning(request, f"No attendance records found for BS {from_bs_date} under the selected filters.")

    return redirect('attendance:dashboard')


@login_required
def daily_report(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    selected_date_raw = request.GET.get('date')
    selected_date, selected_bs_date_obj, selected_bs_date_str = parse_bs_or_ad_date(selected_date_raw)

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')

    selected_grade = None
    selected_section = None
    if selected_grade_id:
        selected_grade = SchoolGrade.objects.filter(id=selected_grade_id).first()
        if selected_section_id:
            selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first()

    all_active_ss = StudentSession.objects.filter(session=current_session, grade__in=grades, status=True).select_related('grade', 'section', 'student')
    all_day_att = DailyAttendance.objects.filter(date=selected_date, session=current_session).select_related('reg_no', 'taken_by', 'grade', 'section')

    if selected_grade:
        all_active_ss = all_active_ss.filter(grade=selected_grade)
        all_day_att = all_day_att.filter(grade=selected_grade)
    if selected_section:
        all_active_ss = all_active_ss.filter(section=selected_section)
        all_day_att = all_day_att.filter(section=selected_section)

    class_summaries = []
    for g in grades:
        g_secs = Section.objects.filter(grade=g)
        if g_secs.exists():
            for sec in g_secs:
                stds = [s for s in all_active_ss if s.grade_id == g.id and s.section_id == sec.id]
                if not stds:
                    continue
                atts = [a for a in all_day_att if a.grade_id == g.id and a.section_id == sec.id]
                class_summaries.append(build_card_dict(g, sec, stds, atts))
        else:
            stds = [s for s in all_active_ss if s.grade_id == g.id]
            if not stds:
                continue
            atts = [a for a in all_day_att if a.grade_id == g.id]
            class_summaries.append(build_card_dict(g, None, stds, atts))

    total_stds = sum(c['total_students'] for c in class_summaries)
    total_present = sum(c['present_total'] for c in class_summaries)
    total_absent = sum(c['absent_total'] for c in class_summaries)
    total_recorded = sum(c['total_students'] for c in class_summaries if c['is_taken'])
    total_pending = total_stds - total_recorded
    pct = round((total_present / total_stds * 100), 1) if total_stds > 0 else 0
    recorded_pct = round((total_recorded / total_stds * 100), 1) if total_stds > 0 else 0

    total_boys = sum(c['total_boys'] for c in class_summaries)
    total_girls = sum(c['total_girls'] for c in class_summaries)
    recorded_boys = sum(c['total_boys'] for c in class_summaries if c['is_taken'])
    recorded_girls = sum(c['total_girls'] for c in class_summaries if c['is_taken'])
    present_boys = sum(c['present_boys'] for c in class_summaries)
    present_girls = sum(c['present_girls'] for c in class_summaries)
    absent_boys = sum(c['absent_boys'] for c in class_summaries)
    absent_girls = sum(c['absent_girls'] for c in class_summaries)
    total_informed_absent = sum(c['informed_absent'] for c in class_summaries)
    total_uninformed_absent = sum(c['uninformed_absent'] for c in class_summaries)
    pending_boys = total_boys - recorded_boys
    pending_girls = total_girls - recorded_girls
    informed_absent_boys = sum(c['informed_absent_boys'] for c in class_summaries)
    informed_absent_girls = sum(c['informed_absent_girls'] for c in class_summaries)
    uninformed_absent_boys = sum(c['uninformed_absent_boys'] for c in class_summaries)
    uninformed_absent_girls = sum(c['uninformed_absent_girls'] for c in class_summaries)

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'selected_grade': selected_grade,
        'selected_section': selected_section,
        'selected_date': selected_date.isoformat(),
        'selected_bs_date_str': selected_bs_date_str,
        'class_summaries': class_summaries,
        'total_stds': total_stds,
        'total_present': total_present,
        'total_absent': total_absent,
        'total_recorded': total_recorded,
        'total_pending': total_pending,
        'percentage': pct,
        'recorded_pct': recorded_pct,
        'total_boys': total_boys,
        'total_girls': total_girls,
        'recorded_boys': recorded_boys,
        'recorded_girls': recorded_girls,
        'present_boys': present_boys,
        'present_girls': present_girls,
        'absent_boys': absent_boys,
        'absent_girls': absent_girls,
        'total_informed_absent': total_informed_absent,
        'total_uninformed_absent': total_uninformed_absent,
        'pending_boys': pending_boys,
        'pending_girls': pending_girls,
        'informed_absent_boys': informed_absent_boys,
        'informed_absent_girls': informed_absent_girls,
        'uninformed_absent_boys': uninformed_absent_boys,
        'uninformed_absent_girls': uninformed_absent_girls,
        'is_report': True,
    }
    return render(request, 'attendance/daily_report.html', context)


@login_required
def print_sheet(request):
    school = get_current_school(request)
    current_session = get_current_session(request)
    
    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')
    selected_date_raw = request.GET.get('date')

    selected_date, selected_bs_date_obj, selected_bs_date_str = parse_bs_or_ad_date(selected_date_raw)

    selected_grade = None
    selected_section = None
    if selected_grade_id:
        selected_grade = grades.filter(id=selected_grade_id).first()
        if not selected_grade:
            messages.error(request, f"Grade not found. The selected grade (ID: {selected_grade_id}) does not exist or is not active.")
            return redirect('attendance:dashboard')
        if selected_section_id:
            selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first()

    ss_query = StudentSession.objects.filter(
        session=current_session,
        status=True
    ).select_related('student', 'section', 'grade')

    if selected_grade:
        ss_query = ss_query.filter(grade=selected_grade)
    if selected_section:
        ss_query = ss_query.filter(section=selected_section)

    student_sessions = list(ss_query.order_by('grade__level__id', 'grade__grade_weight', 'section__section', 'roll_no', 'student__name'))

    att_query = DailyAttendance.objects.filter(
        date=selected_date,
        session=current_session
    )
    if selected_grade:
        att_query = att_query.filter(grade=selected_grade)
    if selected_section:
        att_query = att_query.filter(section=selected_section)

    attendance_map = {}
    taken_by_user = None
    present_count = 0
    absent_count = 0

    for record in att_query:
        attendance_map[record.reg_no_id] = record
        if not taken_by_user and record.taken_by:
            taken_by_user = record.taken_by
        if record.is_present:
            present_count += 1
        else:
            absent_count += 1

    total_stds = len(student_sessions)
    percentage = round((present_count / total_stds) * 100, 1) if total_stds > 0 else 0

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'selected_grade': selected_grade,
        'selected_section': selected_section,
        'selected_date': selected_date.isoformat(),
        'selected_bs_date_str': selected_bs_date_str,
        'student_sessions': student_sessions,
        'attendance_map': attendance_map,
        'taken_by_user': taken_by_user,
        'attendance_summary': {
            'present': present_count,
            'absent': absent_count,
            'percentage': percentage,
        },
        'is_report': True,
    }
    return render(request, 'attendance/print_sheet.html', context)


@login_required
def monthly_report(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')
    
    bs_today = nepali_datetime.date.today()
    
    try:
        selected_year = int(request.GET.get('year', bs_today.year))
    except (ValueError, TypeError):
        selected_year = bs_today.year

    try:
        selected_month = int(request.GET.get('month', bs_today.month))
    except (ValueError, TypeError):
        selected_month = bs_today.month

    month_name = dict(NEPALI_MONTHS).get(selected_month, "")

    selected_grade = None
    selected_section = None
    sections = []

    if selected_grade_id:
        selected_grade = grades.filter(id=selected_grade_id).first()
        if not selected_grade:
            messages.error(request, f"Grade not found. The selected grade (ID: {selected_grade_id}) does not exist or is not active.")
            return redirect('attendance:dashboard')
        sections = Section.objects.filter(grade=selected_grade)
        if selected_section_id:
            selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first()
    elif grades.exists():
        selected_grade = grades.first()
        sections = Section.objects.filter(grade=selected_grade)

    student_rows = []
    total_recorded_days = 0

    if selected_grade:
        ss_query = StudentSession.objects.filter(
            grade=selected_grade,
            session=current_session,
            status=True
        ).select_related('student', 'section')

        if selected_section:
            ss_query = ss_query.filter(section=selected_section)

        student_sessions = list(ss_query.order_by('section__section', 'roll_no', 'student__name'))

        start_ad, end_ad = get_nepali_month_ad_range(selected_year, selected_month)

        att_query = DailyAttendance.objects.filter(
            grade=selected_grade,
            session=current_session,
            date__range=(start_ad, end_ad)
        ).select_related('reg_no', 'taken_by')

        if selected_section:
            att_query = att_query.filter(section=selected_section)

        taken_days = sorted(list(set(att.nepali_date.day for att in att_query if att.nepali_date)))
        total_recorded_days = len(taken_days)

        att_map = {}
        for a in att_query:
            if a.nepali_date:
                reg_no = a.reg_no_id
                day_num = a.nepali_date.day
                if reg_no not in att_map:
                    att_map[reg_no] = {}
                att_map[reg_no][day_num] = a

        for ss in student_sessions:
            student = ss.student
            std_att = att_map.get(student.reg_no, {})
            
            present_cnt = sum(1 for d in std_att.values() if d.is_present)
            absent_cnt = sum(1 for d in std_att.values() if not d.is_present)
            informed_cnt = sum(1 for d in std_att.values() if not d.is_present and d.absent_informed)

            percentage = round((present_cnt / total_recorded_days * 100), 1) if total_recorded_days > 0 else 0

            day_matrix = []
            for d in range(1, 33):
                rec = std_att.get(d)
                if rec:
                    if rec.is_present:
                        day_matrix.append({'status': 'P', 'class': 'table-success font-weight-bold text-success'})
                    elif rec.absent_informed:
                        day_matrix.append({'status': 'L', 'class': 'table-warning font-weight-bold text-warning'})
                    else:
                        day_matrix.append({'status': 'A', 'class': 'table-danger font-weight-bold text-danger'})
                else:
                    day_matrix.append({'status': '-', 'class': 'text-muted'})

            student_rows.append({
                'student_session': ss,
                'student': student,
                'present_cnt': present_cnt,
                'absent_cnt': absent_cnt,
                'informed_cnt': informed_cnt,
                'percentage': percentage,
                'day_matrix': day_matrix,
            })

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'selected_grade': selected_grade,
        'sections': sections,
        'selected_section': selected_section,
        'selected_year': selected_year,
        'selected_month': selected_month,
        'nepali_months': NEPALI_MONTHS,
        'month_name': month_name,
        'years_list': [2085, 2084, 2083, 2082, 2081, 2080],
        'student_rows': student_rows,
        'taken_days_count': total_recorded_days,
        'days_range': list(range(1, 33)),
        'is_report': True,
    }
    return render(request, 'attendance/monthly_report.html', context)


@login_required
def yearly_report(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')

    bs_today = nepali_datetime.date.today()
    
    try:
        selected_year = int(request.GET.get('year', bs_today.year))
    except (ValueError, TypeError):
        selected_year = bs_today.year

    selected_grade = None
    selected_section = None
    sections = []

    if selected_grade_id:
        selected_grade = grades.filter(id=selected_grade_id).first()
        if not selected_grade:
            messages.error(request, f"Grade not found. The selected grade (ID: {selected_grade_id}) does not exist or is not active.")
            return redirect('attendance:dashboard')
        sections = Section.objects.filter(grade=selected_grade)
        if selected_section_id:
            selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first()
    elif grades.exists():
        selected_grade = grades.first()
        sections = Section.objects.filter(grade=selected_grade)

    student_rows = []
    total_recorded_year_days = 0

    if selected_grade:
        ss_query = StudentSession.objects.filter(
            grade=selected_grade,
            session=current_session,
            status=True
        ).select_related('student', 'section')

        if selected_section:
            ss_query = ss_query.filter(section=selected_section)

        student_sessions = list(ss_query.order_by('section__section', 'roll_no', 'student__name'))

        monthly_data = {}
        for m in range(1, 13):
            start_ad, end_ad = get_nepali_month_ad_range(selected_year, m)
            m_att = DailyAttendance.objects.filter(
                grade=selected_grade,
                session=current_session,
                date__range=(start_ad, end_ad)
            )
            if selected_section:
                m_att = m_att.filter(section=selected_section)

            rec_days = len(set(a.date for a in m_att))
            monthly_data[m] = {
                'records': list(m_att),
                'rec_days': rec_days
            }
            total_recorded_year_days += rec_days

        for ss in student_sessions:
            student = ss.student
            months_breakdown = []
            year_present_cnt = 0
            year_absent_cnt = 0

            for m in range(1, 13):
                m_info = monthly_data[m]
                m_recs = [a for a in m_info['records'] if a.reg_no_id == student.reg_no]
                p_cnt = sum(1 for a in m_recs if a.is_present)
                a_cnt = sum(1 for a in m_recs if not a.is_present)
                
                year_present_cnt += p_cnt
                year_absent_cnt += a_cnt

                months_breakdown.append({
                    'month_num': m,
                    'month_name': NEPALI_MONTH_SHORT_NAMES[m-1],
                    'present': p_cnt,
                    'absent': a_cnt,
                    'recorded_days': m_info['rec_days']
                })

            percentage = round((year_present_cnt / total_recorded_year_days * 100), 1) if total_recorded_year_days > 0 else 0

            student_rows.append({
                'student_session': ss,
                'student': student,
                'months_breakdown': months_breakdown,
                'year_present_cnt': year_present_cnt,
                'year_absent_cnt': year_absent_cnt,
                'percentage': percentage
            })

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'selected_grade': selected_grade,
        'sections': sections,
        'selected_section': selected_section,
        'selected_year': selected_year,
        'nepali_month_names': NEPALI_MONTH_SHORT_NAMES,
        'years_list': [2085, 2084, 2083, 2082, 2081, 2080],
        'student_rows': student_rows,
        'total_recorded_year_days': total_recorded_year_days,
        'is_report': True,
    }
    return render(request, 'attendance/yearly_report.html', context)


@login_required
def export_daily_excel(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    selected_date_raw = request.GET.get('date')
    selected_date, selected_bs_date_obj, selected_bs_date_str = parse_bs_or_ad_date(selected_date_raw)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Daily {selected_bs_date_str}"

    title_font = Font(name='Segoe UI', size=14, bold=True, color='008877')
    meta_font = Font(name='Segoe UI', size=10, bold=True)
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='008877', end_color='008877', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0')
    )

    ws.cell(row=1, column=1, value=f"{school.name if school else 'SIMS School'}").font = title_font
    ws.cell(row=2, column=1, value=f"DAILY ATTENDANCE REPORT - BS {selected_bs_date_str} ({selected_date})").font = meta_font

    headers = ['Roll', 'Reg No', 'Student Name', 'Grade', 'Section', 'Status', 'Informed?', 'Remarks', 'Marked By']
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill

    att_records = DailyAttendance.objects.filter(date=selected_date, session=current_session).select_related('reg_no', 'grade', 'section', 'taken_by')

    row_num = 5
    for a in att_records:
        ws.cell(row=row_num, column=1, value='-')
        ws.cell(row=row_num, column=2, value=a.reg_no_id)
        ws.cell(row=row_num, column=3, value=a.reg_no.name if a.reg_no else '')
        ws.cell(row=row_num, column=4, value=a.grade.grade_name if a.grade else '')
        ws.cell(row=row_num, column=5, value=a.section.section if a.section else '-')
        ws.cell(row=row_num, column=6, value='Present' if a.is_present else 'Absent')
        ws.cell(row=row_num, column=7, value='Yes' if a.absent_informed else 'No')
        ws.cell(row=row_num, column=8, value=a.remarks or '')
        ws.cell(row=row_num, column=9, value=a.taken_by.username if a.taken_by else '')

        for col in range(1, 10):
            ws.cell(row=row_num, column=col).border = thin_border
        row_num += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 25
    ws.column_dimensions['I'].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Daily_Attendance_{selected_bs_date_str}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_monthly_excel(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')
    bs_today = nepali_datetime.date.today()
    selected_year = int(request.GET.get('year', bs_today.year))
    selected_month = int(request.GET.get('month', bs_today.month))
    month_name = dict(NEPALI_MONTHS).get(selected_month, "")

    selected_grade = get_object_or_404(SchoolGrade, id=selected_grade_id) if selected_grade_id else SchoolGrade.objects.filter(active=True).first()
    selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first() if (selected_grade and selected_section_id) else None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Monthly {month_name[:4]}"

    title_font = Font(name='Segoe UI', size=14, bold=True, color='008877')
    meta_font = Font(name='Segoe UI', size=10, bold=True)
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='008877', end_color='008877', fill_type='solid')

    present_fill = PatternFill(start_color='E6F4EA', end_color='E6F4EA', fill_type='solid')
    present_font = Font(name='Segoe UI', size=9, bold=True, color='137333')
    absent_fill = PatternFill(start_color='FCE8E6', end_color='FCE8E6', fill_type='solid')
    absent_font = Font(name='Segoe UI', size=9, bold=True, color='C5221F')
    leave_fill = PatternFill(start_color='FEF7E0', end_color='FEF7E0', fill_type='solid')
    leave_font = Font(name='Segoe UI', size=9, bold=True, color='B06000')

    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0')
    )

    ws.cell(row=1, column=1, value=f"{school.name if school else 'SIMS School'}").font = title_font
    ws.cell(row=2, column=1, value=f"MONTHLY ATTENDANCE MATRIX - {month_name.upper()} {selected_year}").font = meta_font
    sec_str = f" - Section {selected_section.section}" if selected_section else ""
    ws.cell(row=3, column=1, value=f"Grade: Grade {selected_grade.grade_name}{sec_str} | Session: {current_session.year}").font = meta_font

    headers = ['Roll', 'Reg No', 'Student Name', 'Present (P)', 'Absent (A)', 'Leave (L)', 'Rate %'] + [str(d) for d in range(1, 33)]
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    start_ad, end_ad = get_nepali_month_ad_range(selected_year, selected_month)
    ss_query = StudentSession.objects.filter(grade=selected_grade, session=current_session, status=True).select_related('student', 'section')
    if selected_section:
        ss_query = ss_query.filter(section=selected_section)
    student_sessions = list(ss_query.order_by('section__section', 'roll_no', 'student__name'))

    att_query = DailyAttendance.objects.filter(grade=selected_grade, session=current_session, date__range=(start_ad, end_ad))
    if selected_section:
        att_query = att_query.filter(section=selected_section)

    taken_days = sorted(list(set(att.nepali_date.day for att in att_query if att.nepali_date)))
    total_recorded_days = len(taken_days)

    att_map = {}
    for a in att_query:
        if a.nepali_date:
            reg = a.reg_no_id
            day_n = a.nepali_date.day
            if reg not in att_map:
                att_map[reg] = {}
            att_map[reg][day_n] = a

    row_num = 6
    for ss in student_sessions:
        student = ss.student
        std_att = att_map.get(student.reg_no, {})
        
        p_cnt = sum(1 for d in std_att.values() if d.is_present)
        a_cnt = sum(1 for d in std_att.values() if not d.is_present)
        l_cnt = sum(1 for d in std_att.values() if not d.is_present and d.absent_informed)
        pct = round((p_cnt / total_recorded_days * 100), 1) if total_recorded_days > 0 else 0

        ws.cell(row=row_num, column=1, value=ss.roll_no).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=student.reg_no).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=3, value=student.name)
        ws.cell(row=row_num, column=4, value=p_cnt).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=5, value=a_cnt).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=6, value=l_cnt).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=7, value=f"{pct}%").alignment = Alignment(horizontal='center')

        for day in range(1, 33):
            col_idx = 7 + day
            cell = ws.cell(row=row_num, column=col_idx)
            cell.alignment = Alignment(horizontal='center')
            rec = std_att.get(day)
            if rec:
                if rec.is_present:
                    cell.value = 'P'
                    cell.fill = present_fill
                    cell.font = present_font
                elif rec.absent_informed:
                    cell.value = 'L'
                    cell.fill = leave_fill
                    cell.font = leave_font
                else:
                    cell.value = 'A'
                    cell.fill = absent_fill
                    cell.font = absent_font
            else:
                cell.value = '-'

        for col in range(1, 40):
            ws.cell(row=row_num, column=col).border = thin_border
        row_num += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 10
    for d in range(1, 33):
        ws.column_dimensions[get_column_letter(7 + d)].width = 4.5

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Monthly_Attendance_{selected_grade.grade_name}_{month_name}_{selected_year}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_yearly_excel(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    selected_grade_id = request.GET.get('grade')
    selected_section_id = request.GET.get('section')
    bs_today = nepali_datetime.date.today()
    selected_year = int(request.GET.get('year', bs_today.year))

    selected_grade = get_object_or_404(SchoolGrade, id=selected_grade_id) if selected_grade_id else SchoolGrade.objects.filter(active=True).first()
    selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first() if (selected_grade and selected_section_id) else None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Yearly {selected_year}"

    title_font = Font(name='Segoe UI', size=14, bold=True, color='008877')
    meta_font = Font(name='Segoe UI', size=10, bold=True)
    header_font = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='008877', end_color='008877', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0')
    )

    ws.cell(row=1, column=1, value=f"{school.name if school else 'SIMS School'}").font = title_font
    ws.cell(row=2, column=1, value=f"YEARLY ATTENDANCE SUMMARY - BS {selected_year}").font = meta_font
    sec_str = f" - Section {selected_section.section}" if selected_section else ""
    ws.cell(row=3, column=1, value=f"Grade: Grade {selected_grade.grade_name}{sec_str} | Session: {current_session.year}").font = meta_font

    ws.merge_cells('A5:A6')
    ws.cell(row=5, column=1, value='Roll').font = header_font
    ws.merge_cells('B5:B6')
    ws.cell(row=5, column=2, value='Reg No').font = header_font
    ws.merge_cells('C5:C6')
    ws.cell(row=5, column=3, value='Student Name').font = header_font

    col_idx = 4
    for m in range(1, 13):
        m_name = NEPALI_MONTH_SHORT_NAMES[m-1]
        c1 = col_idx
        c2 = col_idx + 1
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        cell = ws.cell(row=5, column=c1, value=m_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

        sub1 = ws.cell(row=6, column=c1, value='P')
        sub1.font = header_font
        sub1.fill = header_fill
        sub1.alignment = Alignment(horizontal='center')

        sub2 = ws.cell(row=6, column=c2, value='A')
        sub2.font = header_font
        sub2.fill = header_fill
        sub2.alignment = Alignment(horizontal='center')

        col_idx += 2

    ws.merge_cells(start_row=5, start_column=col_idx, end_row=5, end_column=col_idx+2)
    tot_head = ws.cell(row=5, column=col_idx, value='Year Summary')
    tot_head.font = header_font
    tot_head.fill = header_fill
    tot_head.alignment = Alignment(horizontal='center')

    ws.cell(row=6, column=col_idx, value='Present').font = header_font
    ws.cell(row=6, column=col_idx+1, value='Absent').font = header_font
    ws.cell(row=6, column=col_idx+2, value='Rate %').font = header_font

    for r in range(5, 7):
        for c in range(1, col_idx + 3):
            cell = ws.cell(row=r, column=c)
            cell.fill = header_fill
            cell.border = thin_border

    ss_query = StudentSession.objects.filter(grade=selected_grade, session=current_session, status=True).select_related('student', 'section')
    if selected_section:
        ss_query = ss_query.filter(section=selected_section)
    student_sessions = list(ss_query.order_by('section__section', 'roll_no', 'student__name'))

    monthly_data = {}
    total_recorded_year_days = 0
    for m in range(1, 13):
        start_ad, end_ad = get_nepali_month_ad_range(selected_year, m)
        m_att = DailyAttendance.objects.filter(grade=selected_grade, session=current_session, date__range=(start_ad, end_ad))
        if selected_section:
            m_att = m_att.filter(section=selected_section)
        rec_days = len(set(a.date for a in m_att))
        monthly_data[m] = {'records': list(m_att), 'rec_days': rec_days}
        total_recorded_year_days += rec_days

    row_num = 7
    for ss in student_sessions:
        student = ss.student
        ws.cell(row=row_num, column=1, value=ss.roll_no).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=student.reg_no).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=3, value=student.name)

        c_idx = 4
        y_present = 0
        y_absent = 0

        for m in range(1, 13):
            m_info = monthly_data[m]
            m_recs = [a for a in m_info['records'] if a.reg_no_id == student.reg_no]
            p_cnt = sum(1 for a in m_recs if a.is_present)
            a_cnt = sum(1 for a in m_recs if not a.is_present)

            y_present += p_cnt
            y_absent += a_cnt

            ws.cell(row=row_num, column=c_idx, value=p_cnt).alignment = Alignment(horizontal='center')
            ws.cell(row=row_num, column=c_idx+1, value=a_cnt).alignment = Alignment(horizontal='center')
            c_idx += 2

        pct = round((y_present / total_recorded_year_days * 100), 1) if total_recorded_year_days > 0 else 0

        ws.cell(row=row_num, column=c_idx, value=y_present).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=c_idx+1, value=y_absent).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=c_idx+2, value=f"{pct}%").alignment = Alignment(horizontal='center')

        for c in range(1, c_idx + 3):
            ws.cell(row=row_num, column=c).border = thin_border

        row_num += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 25
    for c in range(4, c_idx + 3):
        ws.column_dimensions[get_column_letter(c)].width = 7

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Yearly_Attendance_{selected_grade.grade_name}_{selected_year}.xlsx"'
    wb.save(response)
    return response


def build_card_dict(grade, section, std_sessions, att_records):
    total_students = len(std_sessions)
    total_boys = sum(1 for s in std_sessions if s.student.gender is True)
    total_girls = sum(1 for s in std_sessions if s.student.gender is False)

    is_taken = len(att_records) > 0
    taken_by = att_records[0].taken_by if is_taken else None
    subject = att_records[0].subject if is_taken else None

    present_total = sum(1 for a in att_records if a.is_present)
    absent_total = sum(1 for a in att_records if not a.is_present)
    informed_absent = sum(1 for a in att_records if not a.is_present and a.absent_informed)
    uninformed_absent = sum(1 for a in att_records if not a.is_present and not a.absent_informed)

    present_boys = sum(1 for a in att_records if a.is_present and a.reg_no.gender is True)
    absent_boys = sum(1 for a in att_records if not a.is_present and a.reg_no.gender is True)
    informed_absent_boys = sum(1 for a in att_records if not a.is_present and a.absent_informed and a.reg_no.gender is True)
    uninformed_absent_boys = sum(1 for a in att_records if not a.is_present and not a.absent_informed and a.reg_no.gender is True)

    present_girls = sum(1 for a in att_records if a.is_present and a.reg_no.gender is False)
    absent_girls = sum(1 for a in att_records if not a.is_present and a.reg_no.gender is False)
    informed_absent_girls = sum(1 for a in att_records if not a.is_present and a.absent_informed and a.reg_no.gender is False)
    uninformed_absent_girls = sum(1 for a in att_records if not a.is_present and not a.absent_informed and a.reg_no.gender is False)

    percentage = round((present_total / total_students) * 100, 1) if (is_taken and total_students > 0) else 0

    return {
        'grade': grade,
        'section': section,
        'total_students': total_students,
        'total_boys': total_boys,
        'total_girls': total_girls,
        'is_taken': is_taken,
        'taken_by': taken_by,
        'subject': subject,
        'present_total': present_total,
        'absent_total': absent_total,
        'informed_absent': informed_absent,
        'uninformed_absent': uninformed_absent,
        'present_boys': present_boys,
        'absent_boys': absent_boys,
        'informed_absent_boys': informed_absent_boys,
        'uninformed_absent_boys': uninformed_absent_boys,
        'present_girls': present_girls,
        'absent_girls': absent_girls,
        'informed_absent_girls': informed_absent_girls,
        'uninformed_absent_girls': uninformed_absent_girls,
        'percentage': percentage,
    }


@login_required
def bulk_backfill_attendance(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    bs_today = nepali_datetime.date.today()
    selected_year = int(request.POST.get('year') or request.GET.get('year') or bs_today.year)
    selected_month = int(request.POST.get('month') or request.GET.get('month') or bs_today.month)

    selected_grade_id = request.GET.get('grade') or request.POST.get('grade')
    selected_section_id = request.GET.get('section') or request.POST.get('section')

    selected_grade = None
    selected_section = None
    sections = []

    if selected_grade_id:
        selected_grade = grades.filter(id=selected_grade_id).first()
        if not selected_grade:
            messages.error(request, f"Grade not found. The selected grade (ID: {selected_grade_id}) does not exist or is not active.")
            return redirect('attendance:dashboard')
        sections = Section.objects.filter(grade=selected_grade)
        if selected_section_id:
            selected_section = Section.objects.filter(grade=selected_grade, id=selected_section_id).first()

    # Calculate valid days in the selected Nepali month
    month_days_all = []
    start_ad = None
    end_ad = None

    WEEKDAY_NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

    for d in range(1, 33):
        try:
            bs_date = nepali_datetime.date(selected_year, selected_month, d)
            ad_date = bs_date.to_datetime_date()
            if not start_ad:
                start_ad = ad_date
            end_ad = ad_date
            month_days_all.append({
                'day': d,
                'bs_date': bs_date,
                'ad_date': ad_date,
                'weekday_name': WEEKDAY_NAMES[ad_date.weekday()],
                'is_saturday': (ad_date.weekday() == 5)
            })
        except Exception:
            break

    # Parse open_days filter
    open_days_raw = request.GET.getlist('open_days') or request.POST.getlist('open_days')
    if open_days_raw:
        selected_open_days = [int(x) for x in open_days_raw if x.isdigit()]
    else:
        # Default to non-Saturday working days
        selected_open_days = [d['day'] for d in month_days_all if not d['is_saturday']]

    # Filtered days list for the matrix columns
    month_days = [d for d in month_days_all if d['day'] in selected_open_days]

    student_sessions = []
    matrix_rows = []

    if selected_grade:
        ss_query = StudentSession.objects.filter(
            grade=selected_grade,
            session=current_session,
            status=True
        ).select_related('student', 'section')

        if selected_section:
            ss_query = ss_query.filter(section=selected_section)

        student_sessions = list(ss_query.order_by('section__section', 'roll_no', 'student__name'))

        # Existing attendance records map for this month
        existing_recs = DailyAttendance.objects.filter(
            grade=selected_grade,
            session=current_session,
            date__range=(start_ad, end_ad),
            subject__isnull=True
        )
        if selected_section:
            existing_recs = existing_recs.filter(section=selected_section)

        existing_map = {}
        for rec in existing_recs:
            try:
                bs_d = nepali_datetime.date.from_datetime(rec.date)
                if bs_d.month == selected_month and bs_d.year == selected_year:
                    existing_map[(rec.reg_no_id, bs_d.day)] = rec
            except Exception:
                pass

        if request.method == 'POST':
            saved_count = 0
            for ss in student_sessions:
                student = ss.student
                for day_info in month_days:
                    d_num = day_info['day']
                    field_name = f"att_{student.reg_no}_{d_num}"
                    status_val = request.POST.get(field_name, '').strip()

                    if status_val in ['P', 'A', 'I']:
                        is_p = (status_val == 'P')
                        is_inf = (status_val == 'I')

                        DailyAttendance.objects.update_or_create(
                            reg_no=student,
                            date=day_info['ad_date'],
                            subject=None,
                            defaults={
                                'school': school or student.school,
                                'session': current_session,
                                'grade': selected_grade,
                                'section': ss.section,
                                'nepali_date': day_info['bs_date'],
                                'is_present': is_p,
                                'absent_informed': is_inf,
                                'remarks': 'Monthly Matrix Backfill',
                                'taken_by': request.user
                            }
                        )
                        saved_count += 1

            month_name = NEPALI_MONTH_SHORT_NAMES[selected_month - 1]
            messages.success(
                request, 
                f"Successfully saved monthly attendance matrix for Grade {selected_grade.grade_name} ({month_name} {selected_year}) across {len(month_days)} open days! ({saved_count} records saved)."
            )
            open_days_query = "&".join(f"open_days={d}" for d in selected_open_days)
            return redirect(f"{request.path}?grade={selected_grade.id}&section={selected_section_id or ''}&year={selected_year}&month={selected_month}&{open_days_query}")

        # Build GET matrix rows
        for ss in student_sessions:
            student = ss.student
            day_statuses = []
            for day_info in month_days:
                d_num = day_info['day']
                rec = existing_map.get((student.reg_no, d_num))
                code = ""
                if rec:
                    if rec.is_present:
                        code = "P"
                    elif rec.absent_informed:
                        code = "I"
                    else:
                        code = "A"
                day_statuses.append({
                    'day': d_num,
                    'code': code,
                    'is_saturday': day_info['is_saturday'],
                    'weekday_name': day_info['weekday_name']
                })
            matrix_rows.append({
                'student_session': ss,
                'student': student,
                'day_statuses': day_statuses
            })

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'selected_grade': selected_grade,
        'sections': sections,
        'selected_section': selected_section,
        'selected_year': selected_year,
        'selected_month': selected_month,
        'nepali_months': NEPALI_MONTHS,
        'years_list': [2085, 2084, 2083, 2082, 2081, 2080],
        'month_days_all': month_days_all,
        'month_days': month_days,
        'selected_open_days': selected_open_days,
        'matrix_rows': matrix_rows,
        'student_sessions': student_sessions,
    }
    return render(request, 'attendance/bulk_backfill.html', context)


@login_required
def student_search(request):
    school = get_current_school(request)
    current_session = get_current_session(request)

    query = request.GET.get('q', '').strip()
    results = []

    if query:
        ss_qs = StudentSession.objects.filter(
            session=current_session,
            student__school=school,
            status=True
        ).select_related('student', 'grade', 'section')

        results_qs = ss_qs.filter(
            Q(student__name__icontains=query) |
            Q(student__reg_no__icontains=query) |
            Q(roll_no__icontains=query)
        )

        total_count = results_qs.count()

        if total_count == 1:
            ss = results_qs.first()
            return redirect('attendance:student_profile', reg_no=ss.student.reg_no)

        paginator = Paginator(results_qs, 50)
        page = request.GET.get('page', 1)
        try:
            results = paginator.page(page)
        except PageNotAnInteger:
            results = paginator.page(1)
        except EmptyPage:
            results = paginator.page(paginator.num_pages)
    else:
        total_count = 0
        results = []

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    context = {
        'school': school,
        'current_session': current_session,
        'grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'nav_grades': grades,
        'query': query,
        'results': results,
        'total_count': total_count,
        'is_report': True,
    }
    return render(request, 'attendance/student_search.html', context)


@login_required
def student_attendance_profile(request, reg_no):
    school = get_current_school(request)
    current_session = get_current_session(request)

    student = Student.objects.filter(reg_no=reg_no).first()
    if not student:
        messages.error(request, f"Student not found with registration number: {reg_no}")
        return redirect('attendance:dashboard')
    student_session = StudentSession.objects.filter(student=student, session=current_session, status=True).select_related('grade', 'section').first()

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    att_records = DailyAttendance.objects.filter(
        reg_no=student,
        session=current_session,
        school=student.school
    ).select_related('taken_by', 'subject').order_by('-date')

    total_records = att_records.count()
    present_count = sum(1 for a in att_records if a.is_present)
    informed_absent_count = sum(1 for a in att_records if not a.is_present and a.absent_informed)
    uninformed_absent_count = sum(1 for a in att_records if not a.is_present and not a.absent_informed)
    absent_count = informed_absent_count + uninformed_absent_count
    percentage = round((present_count / total_records * 100), 1) if total_records > 0 else 0

    bs_today = nepali_datetime.date.today()
    selected_year = int(request.GET.get('year', bs_today.year))
    selected_month = request.GET.get('month', None)
    if selected_month:
        selected_month = int(selected_month)

    monthly_summary = []
    for m in range(1, 13):
        start_ad, end_ad = get_nepali_month_ad_range(selected_year, m)
        m_recs = [a for a in att_records if start_ad <= a.date <= end_ad]
        m_total = len(m_recs)
        m_pres = sum(1 for a in m_recs if a.is_present)
        m_inf = sum(1 for a in m_recs if not a.is_present and a.absent_informed)
        m_uninf = sum(1 for a in m_recs if not a.is_present and not a.absent_informed)
        m_pct = round((m_pres / m_total * 100), 1) if m_total > 0 else 0

        monthly_summary.append({
            'month_num': m,
            'month_name': NEPALI_MONTH_SHORT_NAMES[m - 1],
            'total_recorded': m_total,
            'present': m_pres,
            'informed_absent': m_inf,
            'uninformed_absent': m_uninf,
            'total_absent': m_inf + m_uninf,
            'percentage': m_pct
        })

    # Filter daily records by selected month if provided
    if selected_month:
        month_start_ad, month_end_ad = get_nepali_month_ad_range(selected_year, selected_month)
        filtered_records = [a for a in att_records if month_start_ad <= a.date <= month_end_ad]
        selected_month_name = NEPALI_MONTH_SHORT_NAMES[selected_month - 1]
    else:
        filtered_records = list(att_records[:60])
        selected_month_name = None

    total_present = sum(ms['present'] for ms in monthly_summary)
    total_informed_absent = sum(ms['informed_absent'] for ms in monthly_summary)
    total_uninformed_absent = sum(ms['uninformed_absent'] for ms in monthly_summary)
    total_monthly_records = sum(ms['total_recorded'] for ms in monthly_summary)
    avg_percentage = round((total_present / total_monthly_records * 100), 1) if total_monthly_records > 0 else 0
    audit_logs = DailyAttendance.objects.filter(reg_no=student, session=current_session, audit_logs__isnull=False).select_related('taken_by').prefetch_related('audit_logs').order_by('-date')
    # Only keep audit logs where status actually changed
    from attendance.models import AttendanceAuditLog
    changed_attendance_ids = AttendanceAuditLog.objects.filter(
        attendance__reg_no=student,
        attendance__session=current_session
    ).exclude(previous_status=F('new_status')).values_list('attendance_id', flat=True).distinct()
    audit_logs = DailyAttendance.objects.filter(id__in=changed_attendance_ids).select_related('taken_by').order_by('-date')[:20]

    context = {
        'nav_grades': grades,
        'student': student,
        'student_session': student_session,
        'attendance_records': filtered_records,
        'total_records': total_records,
        'present_count': present_count,
        'absent_count': absent_count,
        'informed_absent_count': informed_absent_count,
        'uninformed_absent_count': uninformed_absent_count,
        'percentage': percentage,
        'selected_year': selected_year,
        'selected_month': selected_month,
        'selected_month_name': selected_month_name,
        'years_list': [2085, 2084, 2083, 2082, 2081, 2080],
        'monthly_summary': monthly_summary,
        'total_present_monthly': total_present,
        'total_informed_absent_monthly': total_informed_absent,
        'total_uninformed_absent_monthly': total_uninformed_absent,
        'avg_percentage': avg_percentage,
        'audit_logs': audit_logs,
        'is_report': True,
    }
    return render(request, 'attendance/student_profile.html', context)


@login_required
def attendance_settings(request):
    """
    Per-school attendance configuration page.
    Creates the settings record on first visit (get_or_create).
    """
    from attendance.models import AttendanceSettings
    from sms.models import SchoolGrade, GradeLevel

    school = get_current_school(request)
    current_session = get_current_session(request)

    if school:
        grades = SchoolGrade.objects.filter(school=school, active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')
    else:
        grades = SchoolGrade.objects.filter(active=True).select_related('level').order_by('level__id', 'grade_weight', 'grade_name')

    nav_level_ids = grades.values_list('level_id', flat=True).distinct()
    nav_grade_levels = list(GradeLevel.objects.filter(id__in=nav_level_ids).order_by('id'))

    if not school:
        messages.error(request, 'No school found for your account.')
        return redirect('attendance:dashboard')

    settings_obj, _ = AttendanceSettings.objects.get_or_create(school=school)

    if request.method == 'POST':
        # Group 1 — subject conflict
        settings_obj.subject_conflict_rule = request.POST.get('subject_conflict_rule', 'absent_wins')

        # Group 2 — notification channels
        settings_obj.notify_sms             = request.POST.get('notify_sms') == 'on'
        settings_obj.notify_push            = request.POST.get('notify_push') == 'on'
        settings_obj.notify_class_group     = request.POST.get('notify_class_group') == 'on'
        settings_obj.notify_student_channel = request.POST.get('notify_student_channel') == 'on'
        settings_obj.notify_student_group   = request.POST.get('notify_student_group') == 'on'
        settings_obj.notify_sims_attendance = request.POST.get('notify_sims_attendance') == 'on'

        # Group 3 — triggers
        settings_obj.notify_on_absent          = request.POST.get('notify_on_absent') == 'on'
        settings_obj.notify_consecutive_only   = request.POST.get('notify_consecutive_only') == 'on'
        settings_obj.consecutive_absence_threshold = int(request.POST.get('consecutive_absence_threshold') or 3)
        settings_obj.notify_on_late_absence    = request.POST.get('notify_on_late_absence') == 'on'
        settings_obj.notify_on_manual_change   = request.POST.get('notify_on_manual_change') == 'on'

        # Group 4 — monthly report
        settings_obj.send_monthly_report = request.POST.get('send_monthly_report') == 'on'

        # Group 5 — locking
        settings_obj.lock_attendance       = request.POST.get('lock_attendance') == 'on'
        settings_obj.lock_after_hours      = int(request.POST.get('lock_after_hours') or 24)
        settings_obj.admin_can_override_lock = request.POST.get('admin_can_override_lock') == 'on'

        settings_obj.updated_by = request.user
        settings_obj.save()

        messages.success(request, 'Attendance settings saved successfully.')
        return redirect('attendance:settings')

    context = {
        'settings': settings_obj,
        'conflict_choices': AttendanceSettings.CONFLICT_CHOICES,
        'nav_grades': grades,
        'nav_grade_levels': nav_grade_levels,
        'is_settings': True,
    }
    return render(request, 'attendance/settings.html', context)
